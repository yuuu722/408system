from pathlib import Path
import json, collections, hashlib, zipfile, re
from pypdf import PdfReader
from docx import Document

root = Path(r'D:\school\科创\平台数据集\组会考研408资料')
out = Path(__file__).parent
files = sorted(p for p in root.rglob('*') if p.is_file())
inventory = [{'path':str(p.relative_to(root)), 'bytes':p.stat().st_size} for p in files]
print('INVENTORY', len(files), dict(collections.Counter(p.suffix.lower() for p in files)), round(sum(p.stat().st_size for p in files)/1024**3,2), flush=True)
sizes=collections.defaultdict(list)
for p in files: sizes[p.stat().st_size].append(p)
hashes=collections.defaultdict(list)
for group in sizes.values():
    if len(group)>1:
        for p in group:
            h=hashlib.file_digest(p.open('rb'),'sha256').hexdigest()
            hashes[h].append(str(p.relative_to(root)))
duplicates=[v for v in hashes.values() if len(v)>1]
samples=[]
for p in files:
    rel=str(p.relative_to(root))
    if p.suffix.lower()=='.docx':
        d=Document(p)
        paras=[x.text for x in d.paragraphs if x.text.strip()]
        row={'path':rel,'type':'docx','paragraphs':len(paras),'tables':len(d.tables),'images':len(d.inline_shapes),'chars':sum(map(len,paras)),'preview':paras[:15]}
        samples.append(row); print(json.dumps(row,ensure_ascii=False),flush=True)
    elif p.suffix.lower()=='.zip':
        with zipfile.ZipFile(p) as z:
            entries=[{'name':x.filename,'bytes':x.file_size} for x in z.infolist() if not x.is_dir()]
            samples.append({'path':rel,'type':'zip','entries':entries})
            print('ZIP', len(entries), flush=True)
    elif p.suffix.lower()=='.pdf':
        chosen = ('原课本' in rel or '27考研' in rel or '2025年计算机408统考真题' in rel or '2024年计算机408统考真题' in rel or '2026年计算机学科' in rel or '考点分类答案' in rel or ('操作系统' in rel and ('速刷' in rel or '真题分类' in rel or '应用题' in rel or '1-109' in rel)))
        if not chosen: continue
        try:
            r=PdfReader(p); n=len(r.pages)
            indices=sorted(set(min(n-1,x) for x in [0, min(5,n-1), n//2, max(0,n-3)]))
            pages=[]
            for i in indices:
                page=r.pages[i]; t=page.extract_text() or ''
                pages.append({'page':i+1,'chars':len(t.strip()),'cjk':len(re.findall(r'[\u4e00-\u9fff]',t)),'preview':t[:1400]})
            row={'path':rel,'type':'pdf','mb':round(p.stat().st_size/1024**2,2),'pages':n,'bookmarks':len(r.outline),'samples':pages}
            samples.append(row)
            print(json.dumps(row,ensure_ascii=False),flush=True)
        except Exception as e:
            samples.append({'path':rel,'error':str(e)})
result={'root':str(root),'inventory':inventory,'duplicates':duplicates,'samples':samples}
(out/'audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print('DUPLICATES',json.dumps(duplicates,ensure_ascii=False),flush=True)
