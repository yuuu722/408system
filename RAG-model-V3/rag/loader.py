"""
文档加载器
将 PDF / Word / TXT / Markdown / PPT 等资料加载为统一的 Document 对象。
解析库采用延迟导入，缺少对应库时给出明确安装提示。
"""
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional


@dataclass
class Document:
    """加载后的文档"""
    text: str
    metadata: dict = field(default_factory=dict)


# 学科识别：根据文件路径中的关键词判断 408 学科
SUBJECT_HINTS = {
    "ds": ["数据结构", "data structure", "ds", "线性表", "树", "图论"],
    "co": ["组成原理", "组成", "计算机组成", "co", "硬件", "指令系统", "cpu"],
    "os": ["操作系统", "os", "进程", "内存管理", "文件系统"],
    "cn": ["网络", "计算机网络", "cn", "tcp", "ip", "协议"],
}


def detect_subject(path: str) -> Optional[str]:
    """根据路径猜测学科"""
    low = str(path).lower()
    for sub, hints in SUBJECT_HINTS.items():
        if any(h in low for h in hints):
            return sub
    return None


def load_file(path: str, subject: Optional[str] = None) -> list[Document]:
    """加载单个文件，返回 Document 列表（一般一个文件一个 Document）"""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(path)

    suffix = p.suffix.lower()
    sub = subject or detect_subject(path)
    meta = {"source": p.name, "path": str(p), "subject": sub or "", "file_type": suffix.lstrip(".")}

    if suffix in (".txt", ".md", ".markdown"):
        text = p.read_text(encoding="utf-8", errors="ignore")
        return [Document(text=text, metadata=meta)]

    if suffix == ".pdf":
        return [Document(text=_load_pdf(p), metadata=meta)]

    if suffix in (".docx", ".doc"):
        return [Document(text=_load_docx(p), metadata=meta)]

    if suffix in (".pptx", ".ppt"):
        return [Document(text=_load_pptx(p), metadata=meta)]

    # 其他类型尝试以文本读取
    try:
        text = p.read_text(encoding="utf-8", errors="ignore")
        return [Document(text=text, metadata=meta)]
    except Exception:
        print(f"[Loader] 暂不支持的文件类型: {suffix} ({p.name})")
        return []


def _load_pdf(path: Path) -> str:
    """加载 PDF，优先 pypdf，其次 pymupdf"""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ImportError:
        pass
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(path))
        return "\n".join(page.get_text() for page in doc)
    except ImportError:
        raise RuntimeError("读取 PDF 需安装解析库：pip install pypdf 或 pip install pymupdf")


def _load_docx(path: Path) -> str:
    """加载 Word"""
    try:
        import docx
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    except ImportError:
        raise RuntimeError("读取 Word 需安装：pip install python-docx")


def _load_pptx(path: Path) -> str:
    """加载 PPT"""
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        return "\n".join(texts)
    except ImportError:
        raise RuntimeError("读取 PPT 需安装：pip install python-pptx")


def load_directory(root: str, subject: Optional[str] = None) -> list[Document]:
    """递归加载整个资料文件夹"""
    root_path = Path(root)
    allowed = {".txt", ".md", ".markdown", ".pdf", ".docx", ".doc", ".pptx", ".ppt"}
    docs = []
    for f in sorted(root_path.rglob("*")):
        if f.is_file() and f.suffix.lower() in allowed:
            # 子文件夹名可能指示学科
            file_subject = subject or detect_subject(str(f))
            try:
                loaded = load_file(str(f), file_subject)
                docs.extend(loaded)
                print(f"[Loader] 已加载 {f.name}（{len(loaded[0].text) if loaded else 0} 字符）")
            except Exception as e:
                print(f"[Loader] 加载失败 {f.name}: {e}")
    return docs
